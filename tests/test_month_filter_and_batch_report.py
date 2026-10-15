from datetime import date
from decimal import Decimal
from pathlib import Path
import openpyxl
import pytest

from src.utils.date_helpers import matches_month, format_periodo_mes, NOMES_MESES
from src.reports.batch_reporter import BatchReporter, EmpresaResumoLote

def test_matches_month_pt_br_names():
    assert matches_month("Setembro", 2026, 9) is True
    assert matches_month("setembro", 2026, 9) is True
    assert matches_month("Set", 2026, 9) is True
    assert matches_month("Janeiro", 2026, 9) is False
    assert matches_month("Fevereiro", 2026, 9) is False
    assert matches_month("Dezembro", 2026, 9) is False

def test_matches_month_rejects_different_year():
    # Deve rejeitar 2025 quando o ano consultado é 2026
    assert matches_month("Setembro/2025", 2026, 9) is False
    assert matches_month("Setembro/2026", 2026, 9) is True
    assert matches_month("2025-09", 2026, 9) is False
    assert matches_month("2026-09", 2026, 9) is True

def test_matches_month_date_formats():
    assert matches_month("15/09/2026", 2026, 9) is True
    assert matches_month("01/09/2026", 2026, 9) is True
    assert matches_month("30/09/2026", 2026, 9) is True
    assert matches_month("15/08/2026", 2026, 9) is False
    assert matches_month("2026-09-15 10:30:00", 2026, 9) is True
    assert matches_month("2026-10-01", 2026, 9) is False

def test_format_periodo_mes():
    d_ini = date(2026, 9, 1)
    d_fim = date(2026, 9, 30)
    fmt = format_periodo_mes(d_ini, d_fim)
    assert "Setembro de 2026" in fmt
    assert "01/09/2026 a 30/09/2026" in fmt

def test_batch_reporter_excel_and_pdf(tmp_path):
    empresas = [
        EmpresaResumoLote(
            razao_social="EMPRESA ALFA LTDA",
            cnpj="63330299000162",
            nfe_saidas_val=Decimal("15000.00"),
            nfe_saidas_qtd=10,
            nfce_saidas_val=Decimal("5000.00"),
            nfce_saidas_qtd=50,
            total_saidas_val=Decimal("20000.00"),
            total_saidas_qtd=60,
            nfe_entradas_val=Decimal("8000.00"),
            nfe_entradas_qtd=5,
            total_canceladas_val=Decimal("0.00"),
            status="Concluído"
        ),
        EmpresaResumoLote(
            razao_social="EMPRESA BETA S.A.",
            cnpj="12720599000117",
            nfe_saidas_val=Decimal("0.00"),
            nfe_saidas_qtd=0,
            nfce_saidas_val=Decimal("0.00"),
            nfce_saidas_qtd=0,
            total_saidas_val=Decimal("0.00"),
            total_saidas_qtd=0,
            nfe_entradas_val=Decimal("0.00"),
            nfe_entradas_qtd=0,
            total_canceladas_val=Decimal("0.00"),
            status="Concluído"
        )
    ]

    excel_path = tmp_path / "relatorio_lote.xlsx"
    pdf_path = tmp_path / "relatorio_lote.pdf"

    ok_xl = BatchReporter.gerar_excel(
        empresas=empresas,
        lote_id=1,
        id_lote_str="test-batch-uuid",
        periodo_str="Setembro de 2026 (01/09/2026 a 30/09/2026)",
        destino_caminho=excel_path
    )
    assert ok_xl is True
    assert excel_path.exists()

    wb = openpyxl.load_workbook(excel_path)
    assert "Relatório Final" in wb.sheetnames
    assert "Consolidado do Lote" in wb.sheetnames
    ws_final = wb["Relatório Final"]
    rows_final = list(ws_final.iter_rows(values_only=True))
    assert rows_final[0] == ("EMPRESAS", "VALOR ENTRADAS", "VALOR SAIDAS")
    assert rows_final[1][0] == "EMPRESA ALFA LTDA"
    assert rows_final[1][1] == 8000.00
    assert rows_final[1][2] == 20000.00
    assert rows_final[2][0] == "EMPRESA BETA S.A."
    assert rows_final[2][1] == 0.00
    assert rows_final[2][2] == 0.00
    assert rows_final[3][0] == "TOTAL GERAL"
    assert rows_final[3][1] == 8000.00
    assert rows_final[3][2] == 20000.00
    ws = wb["Consolidado do Lote"]
    assert "EMPRESA ALFA LTDA" in [str(row[1]) for row in ws.iter_rows(values_only=True)]

    ok_pdf = BatchReporter.gerar_pdf(
        empresas=empresas,
        lote_id=1,
        id_lote_str="test-batch-uuid",
        periodo_str="Setembro de 2026 (01/09/2026 a 30/09/2026)",
        destino_caminho=pdf_path
    )
    assert ok_pdf is True
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 0
